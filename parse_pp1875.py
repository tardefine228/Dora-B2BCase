#!/usr/bin/env python3
"""PDF постановления № 1875 -> pandas DataFrame, CSV, JSON и Excel.

Поддерживаются основные приложения 1–4 с версткой, похожей на исходный PDF.
Номера страниц и количество позиций не зашиты в код. Это не универсальный
парсер произвольных юридических документов. Смотрите warnings.csv.
Python 3.10+. Пример: python parse_pp1875.py pp_1875.pdf --out result
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from statistics import median

import fitz  # пакет PyMuPDF, не пакет с именем fitz
import pandas as pd

VERSION = "1.0.0"
CODE = re.compile(r"(?<![\d.])\d{2}(?:\.\d{1,3}){1,3}(?![\d.])")
POSITION = re.compile(r"^(\d+(?:\.\d+)*)[.)]?$")
FIELDS = [
    "record_id",
    "appendix",
    "position",
    "name",
    "okpd2",
    "code_role",
    "min_share_pct",
    "max_foreign_share_pct",
    "confirming_document",
    "pdf_page_start",
    "pdf_page_end",
    "source_code_cell",
    "review_required",
]


def clean(text: str) -> str:
    """Не удаляем дефисы и условия: они могут менять смысл наименования."""
    return re.sub(r"\s+", " ", text.replace("\u00ad", "").replace("\xa0", " ")).strip()


def word_lines(words: list[dict]) -> list[list[dict]]:
    """Объединение слов по верхней границе строки; координаты от 0 до 1."""
    result = []
    tolerance = max(0.007, median([w["h"] for w in words]) * 0.65) if words else 0.007
    for word in sorted(words, key=lambda w: (w["y"], w["x"])):
        if not result or abs(word["y"] - result[-1][0]["y"]) > tolerance:
            result.append([word])
        else:
            result[-1].append(word)
    return [sorted(line, key=lambda w: w["x"]) for line in result]


def to_text(words: list[dict]) -> str:
    return "\n".join(" ".join(w["t"] for w in line) for line in word_lines(words))


def extract_page(job: tuple) -> dict:
    """Изолированный процесс: PyMuPDF не разделяется между потоками."""
    pdf, index, options, cache_dir = job
    cache = Path(cache_dir) / f"{index + 1:05}.json"
    if cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    with fitz.open(pdf) as doc:
        page = doc[index]
        width, height = page.rect.width, page.rect.height
        native = page.get_text("words", sort=True)
        letters = sum(len(re.findall(r"[А-Яа-яЁё]", w[4])) for w in native)
        words = []
        method = "text"
        if letters >= 80 and not options["force_ocr"]:
            for x0, y0, x1, y1, text, *_ in native:
                words.append(
                    dict(
                        x=x0 / width,
                        y=y0 / height,
                        w=(x1 - x0) / width,
                        h=(y1 - y0) / height,
                        t=text,
                        conf=100.0,
                    )
                )
        else:
            method = "ocr"
            if not options["tesseract"]:
                raise RuntimeError(
                    "Нужен Tesseract с русским языком. Укажите --tesseract или добавьте его в PATH."
                )
            with tempfile.TemporaryDirectory(prefix="pp1875_") as temp:
                image = Path(temp) / "page.png"
                pix = page.get_pixmap(dpi=options["dpi"], alpha=False)
                pix.save(image)
                cmd = [
                    options["tesseract"],
                    str(image),
                    "stdout",
                    "-l",
                    options["lang"],
                    "--psm",
                    "6",
                    "-c",
                    "tessedit_create_tsv=1",
                ]
                if options["tessdata"]:
                    cmd += ["--tessdata-dir", options["tessdata"]]
                run = subprocess.run(
                    cmd,
                    capture_output=True,
                    timeout=options["timeout"],
                    env={**os.environ, "OMP_THREAD_LIMIT": "1"},
                )
                if run.returncode:
                    raise RuntimeError(
                        f"OCR страницы {index + 1}: {run.stderr.decode('utf-8', errors='replace')}"
                    )
                reader = csv.DictReader(
                    io.StringIO(run.stdout.decode("utf-8-sig")),
                    delimiter="\t",
                    quoting=csv.QUOTE_NONE,
                )
                for row in reader:
                    if float(row["conf"]) < 0 or not row.get("text", "").strip():
                        continue
                    words.append(
                        dict(
                            x=int(row["left"]) / pix.width,
                            y=int(row["top"]) / pix.height,
                            w=int(row["width"]) / pix.width,
                            h=int(row["height"]) / pix.height,
                            t=row["text"],
                            conf=float(row["conf"]),
                        )
                    )
    result = {"page": index + 1, "method": method, "words": words, "text": to_text(words)}
    temporary = cache.with_suffix(".tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    temporary.replace(cache)
    return result


def warning(issues: list, code: str, page: int | None, app=None, pos=None, detail=""):
    issues.append(dict(issue=code, pdf_page=page, appendix=app, position=pos, detail=detail))


def appendix_heading(page: dict) -> int | None:
    # Только основные приложения к постановлению, не вложенные формы к Положению.
    top = clean(to_text([w for w in page["words"] if w["y"] < 0.24]))
    match = re.search(r"ПРИЛОЖЕНИЕ\s*[№N]?\s*(\d+)", top, re.I)
    if match and re.search(r"к\s+постановлению\s+Правительства", top, re.I):
        return int(match[1])
    return None


def table_header(words: list[dict], app: int):
    """Нижняя граница повторного заголовка. Не используются номера страниц."""
    if app < 4:
        markers = [w for w in words if "2008" in w["t"] and w["y"] < 0.65]
        has_title = any("Наименование" in w["t"] for w in words if w["y"] < 0.65)
    else:
        markers = [w for w in words if "процентов" in w["t"].lower() and w["y"] < 0.65]
        has_title = any("Подтверждающий" in w["t"] for w in words if w["y"] < 0.65)
    if not markers or not has_title:
        return None
    return min(w["y"] + w["h"] for w in markers) + 0.003


def column_bounds(body: list[dict], all_words: list[dict], app: int):
    """Границы определяются по кодам/числам в текущей таблице, а не пикселям."""
    if app < 4:
        minimum_x = 0.55 if app < 3 else 0.48
        code_words = [w for w in body if w["x"] > minimum_x and CODE.fullmatch(w["t"].rstrip(",;"))]
        if not code_words:
            return None
        left = min(w["x"] for w in code_words) - 0.012
        if app < 3:
            return left, 1.0
        shares = [w for w in body if w["x"] > 0.75 and re.fullmatch(r"\d{1,3}", w["t"])]
        if not shares:
            return None
        return left, min(w["x"] for w in shares) - 0.025
    limits = [w for w in body if 0.32 < w["x"] < 0.65 and w["t"].lower() == "не"]
    docs = [w for w in body if w["x"] > 0.5 and w["t"].lower().startswith("сертификат")]
    if not docs:
        docs = [w for w in all_words if w["x"] > 0.5 and w["t"].startswith("Подтверждающий")]
    if not limits or not docs:
        return None
    return min(w["x"] for w in limits) - 0.02, min(w["x"] for w in docs) - 0.012


def parse_tables(pages: list[dict]) -> tuple[list[dict], list[dict]]:
    positions, issues = [], []
    app = None
    current = None
    seen_apps = set()
    for page in pages:
        number, words = page["page"], page["words"]
        heading = appendix_heading(page)
        if heading is not None:
            app, current = heading, None
            seen_apps.add(app)
            if app not in (1, 2, 3, 4):
                warning(
                    issues,
                    "unsupported_appendix",
                    number,
                    app,
                    detail="Новое приложение: сохранено в тексте PDF, таблица не разобрана.",
                )
        if app not in (1, 2, 3, 4):
            continue
        top = clean(to_text([w for w in words if w["y"] < 0.15])).upper()
        if heading is None and (
            "УТВЕРЖДЕН" in top or "К ПОЛОЖЕНИЮ" in clean(page["text"][:1500]).upper()
        ):
            app, current = None, None
            continue
        cutoff = table_header(words, app)
        if cutoff is None:
            warning(
                issues,
                "unparsed_page",
                number,
                app,
                detail="Не найден заголовок таблицы. Страница сохранена целиком в исходном тексте.",
            )
            current = None  # Не присоединяем следующий фрагмент к неизвестной строке.
            continue
        body = [w for w in words if cutoff < w["y"] < 0.95]
        bounds = column_bounds(body, words, app)
        if bounds is None:
            warning(
                issues,
                "unknown_columns",
                number,
                app,
                detail="Не удалось надёжно определить столбцы. Проверьте исходную страницу.",
            )
            current = None
            continue
        starts = []
        lines = word_lines(body)
        anchors = [
            line[0]["x"]
            for line in lines
            if line[0]["x"] < 0.25 and re.fullmatch(r"\d+(?:\.\d+)*[.)]", line[0]["t"])
        ]
        gutter = median(anchors) + 0.022 if anchors else 0.16
        for line in lines:
            first = line[0]
            match = POSITION.fullmatch(first["t"])
            if match and first["x"] < gutter:
                # Номер позиции должен быть отдельным словом в левом поле.
                starts.append((first, match[1].rstrip(".")))
        starts.sort(key=lambda pair: pair[0]["y"])
        segments = [(None, None, cutoff, starts[0][0]["y"] - 0.004 if starts else 0.95)]
        segments += [
            (w, pos, w["y"] - 0.004, starts[i + 1][0]["y"] - 0.004 if i + 1 < len(starts) else 0.95)
            for i, (w, pos) in enumerate(starts)
        ]
        for start_word, pos, y0, y1 in segments:
            part = [w for w in body if y0 <= w["y"] < y1 and w is not start_word]
            if start_word is not None:
                current = dict(
                    appendix=app,
                    position=pos,
                    names=[],
                    codes=[],
                    shares=[],
                    docs=[],
                    pages=[],
                    low=[],
                )
                positions.append(current)
            if not part:
                continue
            if current is None:
                warning(issues, "orphan_fragment", number, app, detail=clean(to_text(part)))
                continue
            a, b = bounds
            current["names"].append(to_text([w for w in part if w["x"] < a]))
            current["codes" if app < 4 else "shares"].append(
                to_text([w for w in part if a <= w["x"] < b])
            )
            if app == 3:
                current["shares"].append(to_text([w for w in part if w["x"] >= b]))
            if app == 4:
                current["docs"].append(to_text([w for w in part if w["x"] >= b]))
            current["pages"].append(number)
            current["low"] += [
                w["t"] for w in part if w["conf"] < 40 and re.search(r"[А-Яа-яA-Za-z0-9]", w["t"])
            ]
    for missing in sorted({1, 2, 3, 4} - seen_apps):
        warning(
            issues,
            "missing_appendix",
            None,
            missing,
            detail=(
                "Не найдено основное приложение. "
                "Оно могло быть удалено, не входить в PDF или не распознаться."
            ),
        )
    return positions, issues


def make_dataframe(positions: list[dict], issues: list[dict]) -> pd.DataFrame:
    rows, occurrences = [], {}
    for item in positions:
        app, pos = item["appendix"], item["position"]
        key = (app, pos)
        occurrences[key] = occurrences.get(key, 0) + 1
        page = min(item["pages"]) if item["pages"] else None
        last = max(item["pages"]) if item["pages"] else page
        name = clean(" ".join(item["names"]))
        name = re.sub(r"^[—’\s]+", "", name)
        raw_codes = clean(" ".join(item["codes"]))
        share = clean(" ".join(item["shares"]))
        document = clean(" ".join(item["docs"])) or None
        before = len(issues)
        if occurrences[key] > 1:
            warning(
                issues, "duplicate_position", page, app, pos, "Повтор номера: обе записи сохранены."
            )
        if not name:
            warning(issues, "missing_name", page, app, pos, "Не распознано наименование.")
        if item["low"]:
            warning(
                issues, "low_ocr_confidence", page, app, pos, "; ".join(dict.fromkeys(item["low"]))
            )
        if re.search(r"[A-Za-z]", name):
            warning(
                issues,
                "latin_in_name",
                page,
                app,
                pos,
                "Проверьте латинские буквы: это может быть аббревиатура или ошибка OCR.",
            )
        values = re.findall(r"\d+(?:[.,]\d+)?", share)
        percentage = None
        if app in (3, 4):
            if len(values) == 1 and 0 <= float(values[0].replace(",", ".")) <= 100:
                percentage = float(values[0].replace(",", "."))
            else:
                warning(
                    issues, "invalid_percentage", page, app, pos, share or "Доля не распознана."
                )
        matches = list(CODE.finditer(raw_codes)) if app < 4 else []
        if app < 4 and not matches:
            warning(
                issues,
                "missing_code",
                page,
                app,
                pos,
                raw_codes or "Нет кода. Строка сохранена без вымышленного значения.",
            )
        residue = CODE.sub("", raw_codes)
        residue = re.sub(r"за\s+исключением|и", "", residue, flags=re.I)
        if re.search(r"[А-Яа-яA-Za-z0-9]", residue):
            warning(issues, "unparsed_code_text", page, app, pos, raw_codes)
        exception = re.search(r"за\s+исключением", raw_codes, re.I)
        if re.search(r"кроме|исключ", raw_codes, re.I) and not exception:
            warning(issues, "ambiguous_exception", page, app, pos, raw_codes)
        if matches and any(
            raw_codes[m.end() : n.start()].strip() in ("-", "–", "—")
            for m, n in zip(matches, matches[1:], strict=False)
        ):
            warning(
                issues,
                "code_range",
                page,
                app,
                pos,
                "Обнаружен диапазон кодов; границы сохранены, диапазон не развёрнут.",
            )
        for index, match in enumerate(matches or [None], 1):
            role = (
                "Не указан"
                if match is None
                else (
                    "Исключение" if exception and match.start() > exception.start() else "Включён"
                )
            )
            suffix = "" if occurrences[key] == 1 else f"-dup{occurrences[key]}"
            rows.append(
                dict(
                    record_id=f"{app}-{pos}-{index:02}{suffix}",
                    appendix=app,
                    position=pos,
                    name=name,
                    okpd2=match[0] if match else None,
                    code_role=role,
                    min_share_pct=percentage if app == 3 else None,
                    max_foreign_share_pct=percentage if app == 4 else None,
                    confirming_document=document,
                    pdf_page_start=page,
                    pdf_page_end=last,
                    source_code_cell=raw_codes or None,
                    review_required=len(issues) > before,
                )
            )
    for app in sorted({p["appendix"] for p in positions}):
        nums = [
            int(p["position"])
            for p in positions
            if p["appendix"] == app and p["position"].isdigit()
        ]
        if nums:
            gaps = sorted(set(range(1, max(nums) + 1)) - set(nums))
            if gaps:
                warning(
                    issues,
                    "numbering_gap",
                    None,
                    app,
                    detail=(
                        f"Пропуски целых номеров: {gaps}. "
                        "Это может быть изменение документа или ошибка OCR."
                    ),
                )
    df = pd.DataFrame(rows, columns=FIELDS)
    for col in ("appendix", "pdf_page_start", "pdf_page_end"):
        df[col] = df[col].astype("Int64")
    for col in ("record_id", "position", "okpd2"):
        df[col] = df[col].astype("string")
    for col in ("min_share_pct", "max_foreign_share_pct"):
        df[col] = pd.to_numeric(df[col], errors="coerce").astype("Float64")
    df["review_required"] = df["review_required"].astype(bool)
    return df


def export_excel(path: Path, df: pd.DataFrame, issues: list, pages: list, source: str):
    """Полностью Python: openpyxl нужен только для записи Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, headers, records, widths):
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.append(headers)
        for record in records:
            ws.append([None if pd.isna(v) else v for v in record])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Не допускаем интерпретацию текста PDF как формулы Excel.
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243C5A")
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            length = max(
                sum(
                    max(1, (len(line) + int(widths[c]) - 1) // int(widths[c]))
                    for line in str(cell.value or "").split("\n")
                )
                for c, cell in enumerate(row)
            )
            ws.row_dimensions[row[0].row].height = min(409, max(30, length * 14 + 8))
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            table = Table(displayName=f"Table{len(wb.worksheets)}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        return ws

    ws = sheet(
        "Данные",
        FIELDS,
        df.itertuples(index=False, name=None),
        [20, 12, 12, 95, 20, 17, 18, 23, 75, 17, 17, 40, 19],
    )
    for row in ws.iter_rows(min_row=2):
        for index in (0, 2, 4):
            row[index].number_format = "@"
        if row[12].value:
            row[12].fill = PatternFill("solid", fgColor="FFF0CD")
    issue_fields = ["issue", "pdf_page", "appendix", "position", "detail"]
    sheet(
        "Проверить",
        issue_fields,
        ([r.get(k) for k in issue_fields] for r in issues),
        [29, 14, 12, 14, 120],
    )
    source_rows = []
    for page in pages:
        # Дробим текст на короткие блоки, не превышая лимит ячейки Excel.
        lines = page["text"].splitlines()
        for offset in range(0, len(lines), 18):
            source_rows.append(
                [
                    page["page"],
                    offset // 18 + 1,
                    page["method"],
                    "\n".join(lines[offset : offset + 18]),
                ]
            )
    sheet("Текст PDF", ["Страница", "Блок", "Метод", "Текст"], source_rows, [13, 10, 12, 145])
    sheet(
        "Описание",
        ["Параметр", "Значение"],
        [
            ["Источник", source],
            ["Версия парсера", VERSION],
            ["Строк DataFrame", len(df)],
            ["Позиций", len(df[["appendix", "position"]].drop_duplicates())],
            ["Предупреждений", len(issues)],
            ["Единицы долей", "Числа 0–100, например 90 означает 90%. Пусто — не указано."],
            ["Страницы", "Физические страницы PDF, начиная с 1."],
            [
                "Границы",
                "Структурируются основные приложения 1–4. Остальной документ сохранён как текст.",
            ],
            [
                "Проверка",
                "Откройте лист «Проверить». "
                "Отсутствие предупреждений не гарантирует отсутствие ошибок OCR.",
            ],
            ["Исключения", "code_role=Исключение: код не является включённым в позицию."],
            [
                "Редакция",
                "Обработан только входной PDF. Актуальность документа в интернете не проверяется.",
            ],
        ],
        [28, 115],
    )
    wb.save(path)


def run(args) -> int:
    pdf = Path(args.pdf).resolve()
    if not pdf.is_file():
        raise FileNotFoundError(f"PDF не найден: {pdf}")
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    targets = [
        out / name
        for name in (
            "data.csv",
            "data.json",
            "data.xlsx",
            "warnings.csv",
            "report.json",
            "source_text.txt",
        )
    ]
    if not args.overwrite and any(p.exists() for p in targets):
        raise RuntimeError(
            "Результаты уже существуют. Выберите новую папку --out или явно укажите --overwrite."
        )
    executable = args.tesseract or shutil.which("tesseract")
    if not executable and os.name == "nt":
        guess = (
            Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
            / "Tesseract-OCR"
            / "tesseract.exe"
        )
        executable = str(guess) if guess.exists() else None
    engine = "none"
    if executable:
        engine = (
            subprocess.check_output([executable, "--version"], stderr=subprocess.STDOUT)
            .decode(errors="replace")
            .splitlines()[0]
        )
    options = dict(
        tesseract=executable,
        tessdata=str(Path(args.tessdata).resolve()) if args.tessdata else None,
        dpi=args.dpi,
        lang=args.lang,
        force_ocr=args.force_ocr,
        timeout=args.timeout,
    )
    digest = hashlib.sha256()
    with pdf.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    pdf_hash = digest.hexdigest()
    digest.update(json.dumps([VERSION, options, engine], sort_keys=True).encode())
    # Изменение самого traineddata тоже инвалидирует OCR-кеш.
    if args.tessdata:
        for lang in args.lang.split("+"):
            model = Path(args.tessdata) / f"{lang}.traineddata"
            if model.is_file():
                digest.update(model.read_bytes())
    cache = Path(args.cache).resolve() / digest.hexdigest()[:24]
    cache.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf) as doc:
        count = len(doc)
    jobs = [(str(pdf), i, options, str(cache)) for i in range(count)]
    pages = []
    if args.workers == 1:
        results = map(extract_page, jobs)
        for p in results:
            pages.append(p)
            print(f"Страница {len(pages)}/{count}: {p['method']}", flush=True)
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            for p in pool.map(extract_page, jobs):
                pages.append(p)
                print(f"Страница {len(pages)}/{count}: {p['method']}", flush=True)
    positions, issues = parse_tables(pages)
    df = make_dataframe(positions, issues)
    if df.empty:
        warning(
            issues,
            "no_records",
            None,
            detail="Табличные записи не найдены. Проверьте OCR и структуру PDF.",
        )
    df.to_csv(targets[0], sep=";", index=False, encoding="utf-8-sig")
    df.to_json(targets[1], orient="records", force_ascii=False, indent=2)
    export_excel(targets[2], df, issues, pages, pdf.name)
    pd.DataFrame(issues, columns=["issue", "pdf_page", "appendix", "position", "detail"]).to_csv(
        targets[3], sep=";", index=False, encoding="utf-8-sig"
    )
    report = dict(
        parser_version=VERSION,
        source=pdf.name,
        source_sha256=pdf_hash,
        pages=count,
        rows=len(df),
        positions=len(positions),
        warnings=len(issues),
        positions_by_appendix={
            str(k): int(v)
            for k, v in df[["appendix", "position"]]
            .drop_duplicates()
            .groupby("appendix")
            .size()
            .items()
        },
        status="review_required" if issues else "automatic_checks_passed",
    )
    targets[4].write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    targets[5].write_text(
        "\n\n".join(f"=== СТРАНИЦА {p['page']} ===\n{p['text']}" for p in pages), encoding="utf-8"
    )
    print(
        f"Готово: {len(df)} строк; {len(positions)} позиций; "
        f"{len(issues)} предупреждений.\nПапка: {out}"
    )
    return 2 if df.empty or (args.strict and issues) else 0


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("pdf", help="Путь к исходному PDF")
    parser.add_argument("--out", default="result", help="Папка результатов")
    parser.add_argument(
        "--cache", default=".pp1875_cache", help="Кеш распознанных страниц, привязанный к хешу PDF"
    )
    parser.add_argument("--tesseract", help="Путь к tesseract.exe, если он не найден автоматически")
    parser.add_argument("--tessdata", help="Папка с rus.traineddata и eng.traineddata")
    parser.add_argument("--lang", default="rus+eng", help="Языки OCR")
    parser.add_argument("--dpi", type=int, default=240, help="Разрешение OCR (по умолчанию 240)")
    parser.add_argument(
        "--workers", type=int, default=min(4, os.cpu_count() or 1), help="Количество процессов"
    )
    parser.add_argument(
        "--timeout", type=int, default=180, help="Лимит OCR одной страницы в секундах"
    )
    parser.add_argument("--force-ocr", action="store_true", help="Игнорировать текстовый слой PDF")
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Код возврата 2 при предупреждениях; файлы сохраняются",
    )
    parser.add_argument("--overwrite", action="store_true", help="Разрешить перезапись результатов")
    args = parser.parse_args()
    if args.workers < 1 or args.dpi < 100 or args.timeout < 1:
        parser.error("workers >= 1, dpi >= 100, timeout >= 1")
    try:
        return run(args)
    except (OSError, RuntimeError, ValueError, subprocess.SubprocessError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
