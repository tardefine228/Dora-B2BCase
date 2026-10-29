"""Extract product names from an uploaded file (main-page "upload a list" flow).

Supported formats: TXT/MD, CSV/TSV, JSON, XLSX/XLSM, DOCX, PDF.

* Tabular data (CSV, XLSX, DOCX tables, JSON arrays of objects): the column with
  product names is detected by its header ("наименование", "товар", "name", …);
  without a recognizable header the most "text-like" column wins.
* Plain text (TXT, PDF, DOCX paragraphs): one item per line, list markers and
  numbering are stripped.

Heavy parsers (openpyxl, python-docx, PyMuPDF) are imported lazily so the API
still starts without them; a missing parser yields a clear 415-style error.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt": "text",
    ".md": "text",
    ".csv": "csv",
    ".tsv": "csv",
    ".json": "json",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".docx": "docx",
    ".pdf": "pdf",
}

MAX_ITEM_LENGTH = 500
MIN_ITEM_LENGTH = 2

NAME_HEADERS = (
    "наименование",
    "название",
    "товар",
    "продукц",
    "предмет",
    "позиция",
    "номенклатур",
    "описание",
    "окпд",
    "name",
    "product",
    "item",
    "title",
    "description",
)
# Columns that look like names but are not (units, prices, counts…).
NON_NAME_HEADERS = ("ед", "кол", "цена", "сумм", "стоим", "№", "n п/п", "qty", "price", "unit")
JSON_NAME_KEYS = ("name", "title", "product", "item", "наименование", "название", "товар")

_LIST_MARKER = re.compile(r"^\s*(?:[-–—•*·▪●]+|\(?\d{1,4}[.)]|\d{1,4}\s*[-–]\s)\s*")
_ONLY_NUMBERS = re.compile(r"^[\d\s.,;:%№/\\()+\-–—]*$")
_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}(\s|$)")


class FileParseError(ValueError):
    """The file cannot be parsed (bad format, missing parser, no items)."""


class UnsupportedFileError(FileParseError):
    pass


@dataclass
class ExtractedItems:
    format: str
    items: list[str] = field(default_factory=list)
    total_found: int = 0
    duplicates: int = 0
    skipped: int = 0
    truncated: bool = False
    source_column: str | None = None
    warnings: list[str] = field(default_factory=list)


# --- helpers -----------------------------------------------------------------


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def clean_item(raw: object) -> str | None:
    """Normalize one candidate; ``None`` if it is not a product name."""
    if raw is None:
        return None
    text = " ".join(str(raw).replace("\u00a0", " ").split())
    text = _LIST_MARKER.sub("", text).strip(" \t;,")
    if len(text) < MIN_ITEM_LENGTH or _ONLY_NUMBERS.match(text):
        return None
    if not re.search(r"[A-Za-zА-Яа-яЁё]", text):
        return None
    return text[:MAX_ITEM_LENGTH]


def _header_score(header: str) -> int:
    value = header.strip().lower()
    # Our result CSV contains query + title + clause_ref: query is the input.
    if value == "query":
        return len(NAME_HEADERS) + 2
    if not value:
        return 0
    if any(value.startswith(bad) for bad in NON_NAME_HEADERS):
        return -1
    for rank, key in enumerate(NAME_HEADERS):
        if key in value:
            return len(NAME_HEADERS) - rank + 1
    return 0


def _text_score(values: list[str]) -> float:
    """How much a column looks like free-text names (letters, length)."""
    useful = [v for v in values if clean_item(v)]
    if not useful:
        return 0.0
    avg_len = sum(len(v) for v in useful) / len(useful)
    return len(useful) * min(avg_len, 60)


def _pick_from_table(rows: list[list[str]], result: ExtractedItems) -> list[str]:
    rows = [[("" if c is None else str(c)).strip() for c in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    if width == 1:
        body = rows[1:] if rows and _header_score(rows[0][0]) > 0 else rows
        return [row[0] for row in body]

    # Header may not be the first row (titles above the table): look at the top 10.
    best_header = None
    for index, row in enumerate(rows[:10]):
        scores = [_header_score(cell) for cell in row]
        top = max(scores)
        if top > 0 and (best_header is None or top > best_header[2]):
            best_header = (index, scores.index(top), top)

    if best_header:
        header_row, column, _ = best_header
        result.source_column = rows[header_row][column] or None
        return [row[column] for row in rows[header_row + 1 :]]

    columns = [[row[i] for row in rows] for i in range(width)]
    column = max(range(width), key=lambda i: _text_score(columns[i]))
    result.source_column = f"столбец {column + 1}"
    result.warnings.append(
        "Заголовок столбца с наименованиями не найден — выбран столбец "
        f"{column + 1} как наиболее похожий на названия товаров."
    )
    return columns[column]


def _lines(text: str) -> list[str]:
    return text.splitlines()


def _read_text(content: bytes, result: ExtractedItems, filename: str) -> list[str]:
    """Plain text: one item per line.

    Markdown gets one extra rule: a heading (``# Компьютерная техника``,
    ``## Перечень позиций``) is a document structure line, never a product — a
    client's check-list keeps its goods in list items. Without this rule the
    heading is classified as a product and can even come back with a confident
    ``ok`` on a random catalogue position.
    """
    lines = _lines(_decode(content))
    if not filename.lower().endswith((".md", ".markdown")):
        return lines
    kept: list[str] = []
    for line in lines:
        if _MD_HEADING.match(line):
            if line.strip():
                result.skipped += 1
            continue
        kept.append(line)
    if result.skipped:
        result.warnings.append(
            f"Пропущено заголовков Markdown: {result.skipped} — они относятся к структуре "
            "документа, а не к товарам."
        )
    return kept


# --- format readers ----------------------------------------------------------


_CSV_DELIMITERS = (";", ",", "\t", "|")


def _row_widths(sample: str, delimiter: str) -> list[int]:
    """Filled-cell count of every non-empty row parsed with ``delimiter``."""
    widths: list[int] = []
    for row in csv.reader(io.StringIO(sample), delimiter=delimiter):
        filled = sum(1 for cell in row if cell.strip())
        if filled:
            widths.append(filled)
    return widths


def detect_csv_delimiter(text: str) -> str:
    """Pick the delimiter that makes the widest *consistent* table.

    ``csv.Sniffer`` alone misreads the usual Russian exports: a semicolon table
    with decimal commas (``1;Ноутбук;3;89 900,00``) looks comma-separated to it,
    and the whole row becomes one "product name". So candidates are scored by
    the modal column count and how many rows agree with it; the sniffer's answer
    only breaks ties.
    """
    sample = text[:4096]
    try:
        sniffer = csv.Sniffer().sniff(sample, delimiters="".join(_CSV_DELIMITERS))
    except csv.Error:
        sniffed: str | None = None
    else:
        sniffed = sniffer.delimiter

    best: str | None = None
    best_score: tuple[float, float, bool, int] | None = None
    for index, delimiter in enumerate(_CSV_DELIMITERS):
        widths = _row_widths(sample, delimiter)
        if not widths:
            continue
        counts = Counter(widths)
        width, repeats = counts.most_common(1)[0]
        if width < 2:
            continue
        score = (width, repeats / len(widths), delimiter == sniffed, -index)
        if best_score is None or score > best_score:
            best, best_score = delimiter, score
    if best is not None:
        return best
    if sniffed is not None:
        return sniffed
    return ";" if text.count(";") >= text.count(",") else ","


def _read_csv(content: bytes, result: ExtractedItems, filename: str) -> list[str]:
    text = _decode(content)
    delimiter = "\t" if filename.lower().endswith(".tsv") else detect_csv_delimiter(text)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _pick_from_table(rows, result)


def _read_json(content: bytes, result: ExtractedItems) -> list[str]:
    try:
        data = json.loads(_decode(content))
    except json.JSONDecodeError as exc:
        raise FileParseError(f"Некорректный JSON: {exc.msg} (строка {exc.lineno})") from exc
    if isinstance(data, dict):
        for key in ("items", "products", "goods", "data", "rows", "товары"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        raise FileParseError("JSON должен содержать список товаров")
    if all(isinstance(x, dict) for x in data) and data:
        keys: list[str] = []
        for row in data:
            keys.extend(k for k in row if k not in keys)
        for wanted in JSON_NAME_KEYS:
            match = next((k for k in keys if wanted in str(k).lower()), None)
            if match:
                result.source_column = str(match)
                return [row.get(match) for row in data]
        return _pick_from_table([keys] + [[r.get(k) for k in keys] for r in data], result)
    return [x if not isinstance(x, (list, dict)) else None for x in data]


def _read_xlsx(content: bytes, result: ExtractedItems) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UnsupportedFileError("Для XLSX установите пакет openpyxl") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise FileParseError(f"Не удалось открыть Excel-файл: {exc}") from exc
    try:
        best: list[str] = []
        best_meta: tuple[str | None, list[str]] = (None, [])
        for sheet in workbook.worksheets:
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            probe = ExtractedItems(format="xlsx")
            values = _pick_from_table(rows, probe)
            if len([v for v in values if clean_item(v)]) > len([v for v in best if clean_item(v)]):
                best = values
                column = probe.source_column
                best_meta = (f"{sheet.title}: {column}" if column else sheet.title, probe.warnings)
        result.source_column, warnings = best_meta
        result.warnings.extend(warnings)
        return best
    finally:
        workbook.close()


def _read_docx(content: bytes, result: ExtractedItems) -> list[str]:
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedFileError("Для DOCX установите пакет python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise FileParseError(f"Не удалось открыть DOCX: {exc}") from exc
    values: list[str] = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        values.extend(_pick_from_table(rows, result))
    if values:
        return values
    return [p.text for p in document.paragraphs]


def _read_pdf(content: bytes, result: ExtractedItems) -> list[str]:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise UnsupportedFileError("Для PDF установите пакет PyMuPDF") from exc
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise FileParseError(f"Не удалось открыть PDF: {exc}") from exc
    values: list[str] = []
    with document:
        for page in document:
            tables = []
            try:
                tables = page.find_tables().tables
            except Exception:  # older PyMuPDF / no tables
                tables = []
            if tables:
                for table in tables:
                    values.extend(_pick_from_table(table.extract(), result))
            else:
                values.extend(_lines(page.get_text("text")))
    if not values:
        result.warnings.append("В PDF нет текстового слоя (скан?) — распознавание не выполнялось.")
    return values


# --- public API --------------------------------------------------------------


def detect_format(filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileError(
            f"Формат «{suffix or 'без расширения'}» не поддерживается. Допустимо: {allowed}"
        )
    return SUPPORTED_EXTENSIONS[suffix]


def extract_items(filename: str, content: bytes, max_items: int = 100) -> ExtractedItems:
    """Parse ``content`` and return cleaned, de-duplicated product names."""
    kind = detect_format(filename)
    if not content:
        raise FileParseError("Файл пустой")
    result = ExtractedItems(format=kind)
    if kind == "text":
        raw = _read_text(content, result, filename)
    elif kind == "csv":
        raw = _read_csv(content, result, filename)
    elif kind == "json":
        raw = _read_json(content, result)
    elif kind == "xlsx":
        raw = _read_xlsx(content, result)
    elif kind == "docx":
        raw = _read_docx(content, result)
    else:
        raw = _read_pdf(content, result)

    seen: set[str] = set()
    unique: list[str] = []
    for value in raw:
        item = clean_item(value)
        if item is None:
            if value not in (None, "") and str(value).strip():
                result.skipped += 1
            continue
        key = item.casefold()
        if key in seen:
            result.duplicates += 1
            continue
        seen.add(key)
        unique.append(item)

    result.total_found = len(unique)
    if not unique:
        raise FileParseError("В файле не найдено ни одного наименования товара")
    if len(unique) > max_items:
        result.truncated = True
        result.warnings.append(
            f"Найдено {len(unique)} позиций, обработаны первые {max_items} (лимит одной задачи)."
        )
    result.items = unique[:max_items]
    return result
